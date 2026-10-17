"""PDF → Word / Excel dönüşümü (arayüzsüz çekirdek).

* Word: pdf2docx (düzen, tablo, görsel ve yazı biçimi korunur).
* Excel: PyMuPDF `find_tables` ile sayfalardaki tablolar algılanır, openpyxl ile
  her tablo ayrı sayfaya yazılır.

Her iki işlev de bağımsız bir belge kopyası (bayt) alır; bu yüzden arka plan
iş parçacığında ana belgeye dokunmadan çalışır. Çıktı geçici dosyaya yazılıp
`os.replace` ile yerine konur (yarım dosya kalmaz).
"""
from __future__ import annotations

import os
import re
import tempfile
from dataclasses import dataclass, field
from typing import Callable, List, Optional, Sequence

import pymupdf as fitz


class ConvertError(RuntimeError):
    pass


class Cancelled(ConvertError):
    pass


@dataclass
class ConvertResult:
    path: str = ""
    pages: int = 0
    tables: int = 0
    textless_pages: List[int] = field(default_factory=list)     # 0 tabanlı; taranmış olabilir
    notes: List[str] = field(default_factory=list)


def _textless(doc: "fitz.Document", pages: Sequence[int]) -> List[int]:
    return [p for p in pages if len(doc[p].get_text("text").strip()) < 3]


def _atomic_target(out_path: str) -> str:
    folder = os.path.dirname(os.path.abspath(out_path))
    fd, tmp = tempfile.mkstemp(prefix=".donusum_", suffix=os.path.splitext(out_path)[1], dir=folder)
    os.close(fd)
    return tmp


def _check(cancelled: Optional[Callable[[], bool]]):
    if cancelled and cancelled():
        raise Cancelled("Dönüştürme durduruldu.")


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------
def to_docx(pdf_bytes: bytes, out_path: str, pages: Optional[Sequence[int]] = None,
            cancelled: Optional[Callable[[], bool]] = None) -> ConvertResult:
    try:
        from pdf2docx import Converter
    except ImportError as ex:                                   # pragma: no cover
        raise ConvertError("pdf2docx yüklü değil (pip install pdf2docx).") from ex
    with fitz.open("pdf", pdf_bytes) as probe:
        page_list = list(pages) if pages is not None else list(range(probe.page_count))
        res = ConvertResult(pages=len(page_list), textless_pages=_textless(probe, page_list))
    if not page_list:
        raise ConvertError("Dönüştürülecek sayfa yok.")
    _check(cancelled)
    tmp = _atomic_target(out_path)
    cv = Converter(stream=pdf_bytes)
    try:
        cv.convert(tmp, pages=page_list)
    except Exception as ex:
        _silent_remove(tmp)
        raise ConvertError(f"Word'e dönüştürülemedi: {ex}") from ex
    finally:
        cv.close()
    try:
        _check(cancelled)
    except Cancelled:
        _silent_remove(tmp)
        raise
    os.replace(tmp, out_path)
    res.path = out_path
    if res.textless_pages:
        res.notes.append(f"{len(res.textless_pages)} sayfada metin yok (taranmış olabilir); "
                         "bu sayfalar Word'de resim olarak kalır. Önce OCR çalıştırın.")
    return res


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
@dataclass
class TableData:
    page: int                       # 0 tabanlı
    index: int                      # sayfadaki sıra (1'den)
    rows: List[List[str]]


_TR = re.compile(r"^[-+]?\d{1,3}(\.\d{3})+(,\d+)?$|^[-+]?\d+(,\d+)?$")
_EN = re.compile(r"^[-+]?\d{1,3}(,\d{3})+(\.\d+)?$|^[-+]?\d+(\.\d+)?$")


def parse_number(text: str, decimal: str = ","):
    """'1.234,56' (decimal=',') ya da '1,234.56' (decimal='.') → sayı; değilse None.
    Baştaki sıfırlı kodlar ('007'), 15 haneden uzun sayılar ve karışık metinler metin kalır."""
    s = text.strip().replace(" ", "").replace(" ", "")
    if not s or not (_TR if decimal == "," else _EN).match(s):
        return None
    digits = re.sub(r"\D", "", s)
    if len(digits) > 15:
        return None
    body = s.lstrip("+-")
    if len(body) > 1 and body[0] == "0" and body[1] not in ",.":
        return None
    if decimal == ",":
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        f = float(s)
    except ValueError:
        return None
    return int(f) if f.is_integer() and "." not in s else f


def find_tables(pdf_bytes: bytes, pages: Optional[Sequence[int]] = None,
                progress: Optional[Callable[[int, int], None]] = None,
                cancelled: Optional[Callable[[], bool]] = None) -> List[TableData]:
    out: List[TableData] = []
    with fitz.open("pdf", pdf_bytes) as doc:
        page_list = list(pages) if pages is not None else list(range(doc.page_count))
        for n, p in enumerate(page_list):
            _check(cancelled)
            try:
                found = doc[p].find_tables()
            except Exception:                                     # bozuk sayfa diğerlerini durdurmasın
                found = None
            idx = 0
            for t in (found.tables if found else []):
                rows = [[("" if c is None else " ".join(str(c).split())) for c in r]
                        for r in t.extract()]
                rows = [r for r in rows if any(r)]
                if len(rows) < 2 or max((len(r) for r in rows), default=0) < 2:
                    continue
                idx += 1
                out.append(TableData(p, idx, rows))
            if progress:
                progress(n + 1, len(page_list))
    return out


def write_xlsx(tables: List[TableData], out_path: str, numbers: bool = True,
               decimal: str = ",") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, Side
        from openpyxl.utils import get_column_letter
    except ImportError as ex:                                   # pragma: no cover
        raise ConvertError("openpyxl yüklü değil (pip install openpyxl).") from ex
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for t in tables:
        name = "Tablo" if len(tables) == 1 else f"S{t.page + 1}-T{t.index}"
        ws = wb.create_sheet(name)
        width = max(len(r) for r in t.rows)
        for r_i, row in enumerate(t.rows, start=1):
            for c_i in range(width):
                text = row[c_i] if c_i < len(row) else ""
                value = parse_number(text, decimal) if numbers and r_i > 1 else None
                cell = ws.cell(row=r_i, column=c_i + 1, value=text if value is None else value)
                cell.border = border
                cell.alignment = Alignment(wrap_text=True, vertical="top",
                                           horizontal="right" if value is not None else None)
        for c in ws[1]:
            c.font = Font(bold=True)
        for c_i in range(1, width + 1):
            longest = max((len(str(ws.cell(row=r, column=c_i).value or ""))
                           for r in range(1, len(t.rows) + 1)), default=8)
            ws.column_dimensions[get_column_letter(c_i)].width = min(max(longest + 2, 8), 60)
    tmp = _atomic_target(out_path)
    try:
        wb.save(tmp)
    except Exception:
        _silent_remove(tmp)
        raise
    os.replace(tmp, out_path)


def to_xlsx(pdf_bytes: bytes, out_path: str, pages: Optional[Sequence[int]] = None,
            numbers: bool = True, decimal: str = ",",
            progress: Optional[Callable[[int, int], None]] = None,
            cancelled: Optional[Callable[[], bool]] = None) -> ConvertResult:
    with fitz.open("pdf", pdf_bytes) as probe:
        page_list = list(pages) if pages is not None else list(range(probe.page_count))
        res = ConvertResult(pages=len(page_list), textless_pages=_textless(probe, page_list))
    tables = find_tables(pdf_bytes, page_list, progress, cancelled)
    res.tables = len(tables)
    if not tables:
        hint = (" Sayfalar taranmış görünüyor; önce OCR çalıştırın."
                if len(res.textless_pages) == len(page_list) else "")
        raise ConvertError("Seçilen sayfalarda tablo algılanamadı." + hint)
    _check(cancelled)
    write_xlsx(tables, out_path, numbers, decimal)
    res.path = out_path
    if res.textless_pages:
        res.notes.append(f"{len(res.textless_pages)} sayfada metin yok; tablolar bu sayfalardan alınamaz (OCR gerekir).")
    return res


def _silent_remove(path: str):
    try:
        os.remove(path)
    except OSError:
        pass
