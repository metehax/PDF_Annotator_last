import os, sys, tempfile
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "src"))
os.chdir(ROOT)
import pymupdf as fitz

tmp = tempfile.mkdtemp(prefix="d2_")
ok = []


def check(name, cond, extra=""):
    ok.append(bool(cond))
    print(("OK   " if cond else "FAIL ") + name + (f"  [{extra}]" if extra and not cond else ""))


def make_doc(path):
    d = fitz.open()
    for pg in range(5):
        p = d.new_page()
        p.insert_text((72, 80), "Yillik Rapor 2025" if pg == 0 else f"Bolum {pg}", fontsize=24)
        p.insert_text((72, 120), "Bu belge icin bir ornek metin yazilmistir ve bu bir deneme icindir.", fontsize=11)
        p.insert_text((72, 140), "Taraflar ile ilgili ayrintilar asagida bir tabloda verilmistir.", fontsize=11)
        if pg == 1:
            rows = [["Kalem", "Adet", "Tutar"], ["Kalem A", "10", "1.250,50"], ["Kalem B", "7", "980"], ["Kalem C", "007", "12,5"]]
            x0, y0, cw, ch = 72, 200, 120, 24
            for r, row in enumerate(rows):
                for c, txt in enumerate(row):
                    rect = fitz.Rect(x0 + c * cw, y0 + r * ch, x0 + (c + 1) * cw, y0 + (r + 1) * ch)
                    p.draw_rect(rect, color=(0, 0, 0), width=0.8)
                    p.insert_text((rect.x0 + 5, rect.y1 - 8), txt, fontsize=11)
        if pg == 2:
            pm = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 100, 100), False)
            pm.set_rect(pm.irect, (200, 30, 30))
            p.insert_image(fitz.Rect(72, 200, 172, 300), pixmap=pm)
    d.save(path)
    d.close()


SRC = os.path.join(tmp, "rapor.pdf")
make_doc(SRC)
data = open(SRC, "rb").read()

import convert
# ---- Word
out = os.path.join(tmp, "rapor.docx")
res = convert.to_docx(data, out)
check("docx yazildi", os.path.getsize(out) > 5000 and res.pages == 5)
import docx
dd = docx.Document(out)
txt = "\n".join(p.text for p in dd.paragraphs)
check("docx metni var", "Yillik Rapor" in txt, txt[:80])
check("docx tablo var", len(dd.tables) >= 1, len(dd.tables))
r1 = convert.to_docx(data, os.path.join(tmp, "tek.docx"), pages=[1])
check("tek sayfa dokumani", r1.pages == 1 and os.path.exists(os.path.join(tmp, "tek.docx")))
try:
    convert.to_docx(data, os.path.join(tmp, "iptal.docx"), cancelled=lambda: True)
    check("iptal firlatir", False)
except convert.Cancelled:
    check("iptal firlatir", not os.path.exists(os.path.join(tmp, "iptal.docx")))
check("gecici dosya kalmadi", not [f for f in os.listdir(tmp) if f.startswith(".donusum_")], os.listdir(tmp))

# ---- Excel
xo = os.path.join(tmp, "rapor.xlsx")
xr = convert.to_xlsx(data, xo)
check("xlsx yazildi, 1 tablo", xr.tables == 1 and os.path.exists(xo), xr.tables)
import openpyxl
wb = openpyxl.load_workbook(xo)
ws = wb.worksheets[0]
vals = [[c.value for c in row] for row in ws.iter_rows()]
print("xlsx:", vals)
check("sayfa adi", ws.title == "Tablo", ws.title)
check("baslik satiri metin", vals[0] == ["Kalem", "Adet", "Tutar"], vals[0])
check("1.250,50 -> 1250.5", vals[1][2] == 1250.5, vals[1])
check("10 -> int", vals[1][1] == 10 and isinstance(vals[1][1], int), vals[1])
check("007 metin kalir", vals[3][1] == "007", vals[3])
check("12,5 -> 12.5", vals[3][2] == 12.5, vals[3])
xo2 = os.path.join(tmp, "metin.xlsx")
convert.to_xlsx(data, xo2, numbers=False)
check("numbers=False metin", openpyxl.load_workbook(xo2).worksheets[0]["C2"].value == "1.250,50")
try:
    convert.to_xlsx(data, os.path.join(tmp, "x.xlsx"), pages=[0])
    check("tablosuz sayfa hata", False)
except convert.ConvertError as ex:
    check("tablosuz sayfa hata", "tablo" in str(ex).lower())
check("parse_number en", convert.parse_number("1,234.5", ".") == 1234.5)
check("parse_number karisik metin", convert.parse_number("12 adet") is None and convert.parse_number("") is None)

# ---- Erisilebilirlik
import a11y
doc = fitz.open(SRC)
rep = a11y.audit(doc, SRC)
st = {c.key: c.status for c in rep.checks}
print("durum:", st)
check("etiketsiz -> fail", st["tagged"] == "fail")
check("dil yok -> fail + oneri tr-TR", st["lang"] == "fail" and rep.get("lang").data["suggest"] == "tr-TR", rep.get("lang").data)
check("baslik bos -> fail + oneri", st["title"] == "fail" and "Yillik" in rep.get("title").data["suggest"], rep.get("title").data)
check("baslik yapisi -> yer imi onerisi", st["headings"] == "warn" and rep.get("headings").fix == "bookmarks", (st["headings"], rep.get("headings").detail))
check("alt metin -> fail (gorsel)", st["alt"] == "fail", rep.get("alt").detail)
check("yazi tipi gomulu degil -> warn", st["fonts"] == "warn", rep.get("fonts").detail)
check("taranmis yok -> pass", st["scanned"] == "pass")
check("okuma sirasi pass", st["order"] == "pass", rep.get("order").detail)
check("duzeltilebilir 3", {c.fix for c in rep.fixable()} == {"lang", "title", "bookmarks"}, [c.fix for c in rep.fixable()])
toc = rep.get("headings").data["toc"]
check("toc ilk duzey 1", toc and toc[0][0] == 1 and all(b[0] - a[0] <= 1 for a, b in zip(toc, toc[1:])), toc[:3])
done = a11y.apply_fixes(doc, ["lang", "title", "bookmarks"], lang="tr-TR",
                        title=rep.get("title").data["suggest"], toc=toc)
print("duzeltildi:", done)
rep2 = a11y.audit(doc, SRC)
st2 = {c.key: c.status for c in rep2.checks}
check("duzeltme sonrasi lang/title/headings gecer", st2["lang"] == st2["title"] == st2["headings"] == "pass", st2)
# kaydet-yukle: kalicilik
FIXED = os.path.join(tmp, "fixed.pdf")
doc.save(FIXED)
d3 = fitz.open(FIXED)
st3 = {c.key: c.status for c in a11y.audit(d3, FIXED).checks}
check("kayit sonrasi kalici", st3["lang"] == st3["title"] == st3["headings"] == "pass", st3)
# dosya adi gibi baslik
d3.set_metadata({"title": "Microsoft Word - rapor.docx"})
check("dosya adi basligi uyarilir", a11y.audit(d3).get("title").status == "fail")

# taranmis sayfa + OCR
import ocr
src2 = fitz.open(SRC)
pix = src2[0].get_pixmap(dpi=150, alpha=False)
sc = fitz.open()
for _ in range(2):
    pg = sc.new_page(width=src2[0].rect.width, height=src2[0].rect.height)
    pg.insert_image(pg.rect, pixmap=pix)
srep = a11y.audit(sc)
check("taranmis sayfalar bulundu", srep.get("scanned").status == "fail" and srep.get("scanned").data["pages"] == [0, 1], srep.get("scanned").data)
check("taranmis: alt metin sayilmaz", srep.get("alt").status == "pass", srep.get("alt").detail)
tess = ocr.find_tessdata()
print("tessdata:", tess)
if tess:
    lang = ocr.pick_language(tess, "eng")
    results = []
    for i in (0, 1):
        pm = sc[i].get_pixmap(dpi=200, alpha=False)
        results.append((i, pm.pdfocr_tobytes(language=lang, tessdata=tess)))
    n = a11y.ocr_replace_pages(sc, results)
    check("ocr sayfa degisti", n == 2 and sc.page_count == 2)
    check("sayfa boyutu ayni", abs(sc[0].rect.width - src2[0].rect.width) < 0.5 and abs(sc[0].rect.height - src2[0].rect.height) < 0.5)
    check("ocr metni cikarilabiliyor", "rapor" in sc[0].get_text().lower() or "yillik" in sc[0].get_text().lower(), sc[0].get_text()[:60])
    check("ocr sonrasi taranmis yok", a11y.audit(sc).get("scanned").status == "pass")
else:
    print("OCR atlandi (tessdata yok)")
# rapor ciktilari
check("rapor txt/html", "Erisilebilirlik" in a11y.report_text(rep).replace("ş", "s").replace("Erişilebilirlik", "Erisilebilirlik")
      and "<html" in a11y.report_html(rep))
# etiketli belge
tg = fitz.open(SRC)
cat = tg.pdf_catalog()
root = tg.get_new_xref()
tg.update_object(root, "<</Type/StructTreeRoot>>")
tg.xref_set_key(cat, "StructTreeRoot", f"{root} 0 R")
tg.xref_set_key(cat, "MarkInfo", "<</Marked true>>")
fig = tg.get_new_xref()
tg.update_object(fig, "<</Type/StructElem/S/Figure/P %d 0 R>>" % root)
trep = a11y.audit(tg)
check("etiketli: tagged pass", trep.get("tagged").status == "pass")
check("etiketli: alt metinsiz sekil fail", trep.get("alt").status == "fail", trep.get("alt").detail)
tg.xref_set_key(fig, "Alt", "(Kirmizi kare)")
check("etiketli: alt metin var", a11y.audit(tg).get("alt").status == "pass")
check("etiketli: okuma sirasi info", trep.get("order").status == "info")

print(f"\nOZET: {sum(ok)}/{len(ok)} gecti")
sys.stdout.flush()
os._exit(0 if all(ok) else 1)
